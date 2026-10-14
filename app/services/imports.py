"""Bulk CSV/XLSX import into dynamic tables, executed by a background worker.

Validation happens twice:

* **On upload** (synchronous, cheap). Covers size, format, zip-bomb limits for XLSX, the header row
  against the table's fields, and the uploader's create permission and field-level write permission.
* **In the worker**, for every row. Values are typed with the same coercion as the API, and the
  uploader's permissions are re-evaluated *at execution time*.

Idempotency and retries: each row is inserted with a deterministic id ``uuid5(task_id, row_number)``
using ``ON CONFLICT DO NOTHING``. A retried or duplicated import therefore never creates duplicates,
for shared tenants and for dedicated databases alike. The import runs in a single transaction, so
``all_or_nothing`` (the default) means exactly that. In ``skip_invalid`` mode the valid rows are
inserted and the invalid ones are reported.
"""

from __future__ import annotations

import csv
import datetime as dt
import hashlib
import io
import uuid
import zipfile
from collections.abc import Iterator
from datetime import UTC, datetime, timedelta
from typing import Any, Literal

from sqlalchemy import select
from sqlalchemy.dialects.postgresql import insert as pg_insert

from app.core.config import Settings
from app.core.errors import AppError, Forbidden, NotFound, ValidationFailed
from app.core.logging import get_logger
from app.db.session import Database
from app.dynamic.query import coerce
from app.dynamic.schema import TableMeta
from app.models import Action, CrudMode, IsolationMode, SystemFile, SystemTask, TaskStatus
from app.services.audit import AuditContext, record
from app.services.authz import AuthzService, Principal
from app.services.outbox import emit
from app.services.schema import SchemaService
from app.services.tenant_db import TenantDbService
from app.services.tenants import TenantService

log = get_logger("imports")
Mode = Literal["all_or_nothing", "skip_invalid"]
Fmt = Literal["csv", "xlsx"]
MAX_ERRORS = 200
XLSX_MAX_ENTRIES = 2000


class ImportFormatError(ValidationFailed):
    pass


def detect_format(data: bytes, filename: str) -> Fmt:
    if data[:4] == b"PK\x03\x04":
        if not filename.lower().endswith(".xlsx"):
            raise ImportFormatError("Zip-based uploads must be .xlsx workbooks.")
        return "xlsx"
    if not filename.lower().endswith(".csv"):
        raise ImportFormatError("Only .csv and .xlsx files are accepted.")
    return "csv"


def _check_xlsx(data: bytes, max_uncompressed: int) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as zf:
            infos = zf.infolist()
            if len(infos) > XLSX_MAX_ENTRIES:
                raise ImportFormatError("Workbook has too many parts.")
            total = sum(i.file_size for i in infos)
            if total > max_uncompressed:
                raise ImportFormatError("Workbook expands beyond the allowed size.")
            names = {i.filename for i in infos}
            if "[Content_Types].xml" not in names or not any(n.startswith("xl/worksheets/") for n in names):
                raise ImportFormatError("Not a valid .xlsx workbook.")
    except zipfile.BadZipFile as e:
        raise ImportFormatError("Corrupt .xlsx file.") from e


def iter_rows(data: bytes, fmt: Fmt, max_uncompressed: int) -> Iterator[list[Any]]:
    """Yield rows as lists (header first). Empty cells become None."""
    if fmt == "csv":
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError as e:
            raise ImportFormatError("CSV must be UTF-8 encoded.") from e
        if "\x00" in text:
            raise ImportFormatError("CSV contains NUL bytes.")
        csv.field_size_limit(1_000_000)
        for row in csv.reader(io.StringIO(text, newline="")):
            yield [c if c != "" else None for c in row]
        return
    _check_xlsx(data, max_uncompressed)
    from openpyxl import load_workbook  # heavy import, worker/upload only

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        ws = wb.worksheets[0]
        for xrow in ws.iter_rows(values_only=True):
            yield [None if (v is None or v == "") else v for v in xrow]
    finally:
        wb.close()


def _normalise(v: Any) -> Any:
    # Spreadsheet datetimes are naive: interpret them as UTC (documented).
    if isinstance(v, dt.datetime) and v.tzinfo is None:
        return v.replace(tzinfo=UTC)
    return v


def validate_header(meta: TableMeta, header: list[Any], writable: frozenset[str] | None) -> list[str]:
    names = [str(h).strip() if h is not None else "" for h in header]
    while names and names[-1] == "":
        names.pop()
    fmap = meta.field_map
    unknown = sorted({n for n in names if n not in fmap})
    if unknown or not names:
        raise ValidationFailed(
            "Header contains unknown or system columns.", details={"columns": unknown[:20]}
        )
    if len(set(names)) != len(names):
        raise ValidationFailed("Header contains duplicate columns.")
    if writable is not None and (forbidden := sorted(set(names) - writable)):
        raise Forbidden("You may not write some of the columns in this file.", details={"columns": forbidden})
    missing = sorted(f.name for f in meta.fields if not f.nullable and f.name not in names)
    if missing:
        raise ValidationFailed("Required columns are missing.", details={"columns": missing})
    return names


class ImportService:
    def __init__(
        self,
        settings: Settings,
        db: Database,
        schema: SchemaService,
        authz: AuthzService,
        tenants: TenantService,
        tenant_db: TenantDbService,
    ) -> None:
        self._s = settings
        self._db = db
        self._schema = schema
        self._authz = authz
        self._tenants = tenants
        self._tenant_db = tenant_db

    async def _meta(self, table: str, p: Principal) -> TableMeta:
        meta = await self._schema.get_table(table)
        if meta is None or meta.crud_mode == CrudMode.disabled:
            raise NotFound()
        if meta.crud_mode == CrudMode.read_only:
            raise Forbidden("This table is read-only for generic writes.")
        p.require("imports", Action.create)
        p.require(meta.resource, Action.create)
        return meta

    async def submit(
        self,
        p: Principal,
        table: str,
        *,
        data: bytes,
        filename: str,
        mode: Mode,
        idempotency_key: str | None,
        ctx: AuditContext,
    ) -> tuple[SystemTask, bool]:
        """Validate cheaply, persist the file and a task row in one transaction. Returns (task, created)."""
        meta = await self._meta(table, p)
        if not data:
            raise ValidationFailed("Empty file.")
        if len(data) > self._s.import_max_bytes:
            raise ValidationFailed(f"File exceeds {self._s.import_max_bytes} bytes.")
        safe_name = "".join(ch for ch in filename if ch.isalnum() or ch in "._- ")[:200] or "upload"
        fmt = detect_format(data, safe_name)
        rows = iter_rows(data, fmt, self._s.import_max_uncompressed_bytes)
        header = next(rows, None)
        if header is None:
            raise ValidationFailed("File has no header row.")
        grant = p.require(meta.resource, Action.create)
        validate_header(meta, header, grant.writable)

        async with self._db.tenant_session(p.tenant_id, user_id=p.user_id) as s:
            if idempotency_key:
                existing = (
                    await s.execute(
                        select(SystemTask).where(
                            SystemTask.kind == "import", SystemTask.idempotency_key == idempotency_key
                        )
                    )
                ).scalar_one_or_none()
                if existing is not None:
                    return existing, False
            file = SystemFile(
                tenant_id=p.tenant_id,
                purpose="import",
                filename=safe_name,
                content_type="text/csv"
                if fmt == "csv"
                else "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                size_bytes=len(data),
                sha256=hashlib.sha256(data).hexdigest(),
                content=data,
                created_by=p.user_id,
                expires_at=datetime.now(UTC) + timedelta(days=self._s.file_retention_days),
            )
            s.add(file)
            await s.flush()
            task = SystemTask(
                tenant_id=p.tenant_id,
                kind="import",
                idempotency_key=idempotency_key,
                created_by=p.user_id,
                input={
                    "file_id": str(file.id),
                    "table": meta.name,
                    "mode": mode,
                    "format": fmt,
                    "actor_kind": p.kind,
                    "actor_id": str(p.actor_id),
                    "scopes": sorted(p.scopes) if p.scopes else None,
                },
            )
            s.add(task)
            await s.flush()
            await record(
                s,
                ctx,
                tenant_id=p.tenant_id,
                action="import.submit",
                resource=meta.resource,
                resource_id=str(task.id),
                detail={"file": safe_name, "bytes": len(data), "mode": mode},
            )
            await s.commit()
        return task, True

    async def run(self, tenant_id: uuid.UUID, task_id: uuid.UUID) -> TaskStatus:
        """Worker entry point. Safe to call repeatedly (retries, duplicate messages)."""
        async with self._db.tenant_session(tenant_id) as s:
            task = (
                await s.execute(select(SystemTask).where(SystemTask.id == task_id).with_for_update())
            ).scalar_one_or_none()
            if task is None:
                return TaskStatus.failed
            if task.status in (TaskStatus.succeeded, TaskStatus.failed):
                return task.status
            task.status, task.attempts = TaskStatus.running, task.attempts + 1
            task.started_at = task.started_at or datetime.now(UTC)
            inp, created_by = dict(task.input), task.created_by
            file = (
                await s.execute(select(SystemFile).where(SystemFile.id == uuid.UUID(inp["file_id"])))
            ).scalar_one_or_none()
            data = file.content if file else None
            await s.commit()
        try:
            if data is None or created_by is None:
                raise ValidationFailed("The uploaded file is no longer available.")
            result = await self._execute(tenant_id, task_id, created_by, inp, data)
        except AppError as e:
            await self._close(tenant_id, task_id, TaskStatus.failed, None, e.message, e.details)
            return TaskStatus.failed
        await self._close(tenant_id, task_id, TaskStatus.succeeded, result, None, None)
        return TaskStatus.succeeded

    async def _execute(
        self, tenant_id: uuid.UUID, task_id: uuid.UUID, user_id: uuid.UUID, inp: dict[str, Any], data: bytes
    ) -> dict[str, Any]:
        # Re-check authorisation as of now: the uploader may have lost access since submitting.
        matrix = await self._authz.matrix(tenant_id, user_id)
        scopes = frozenset(inp["scopes"]) if inp.get("scopes") else None
        p = Principal(
            inp.get("actor_kind", "user"),
            tenant_id,
            user_id,
            uuid.UUID(inp["actor_id"]),
            None,
            matrix,
            scopes,
        )
        meta = await self._meta(inp["table"], p)
        grant = p.require(meta.resource, Action.create)
        rows = iter_rows(data, inp["format"], self._s.import_max_uncompressed_bytes)
        names = validate_header(meta, next(rows, []), grant.writable)
        fmap = meta.field_map

        valid: list[dict[str, Any]] = []
        errors: list[dict[str, Any]] = []
        invalid = 0
        for n, raw in enumerate(rows, start=2):
            if n - 1 > self._s.import_max_rows:
                raise ValidationFailed(f"File exceeds {self._s.import_max_rows} data rows.")
            if all(v is None for v in raw):
                continue
            try:
                values = {
                    name: coerce(fmap[name], _normalise(raw[i]) if i < len(raw) else None, from_query=True)
                    for i, name in enumerate(names)
                }
                missing = [f.name for f in meta.fields if not f.nullable and values.get(f.name) is None]
                if missing:
                    raise ValidationFailed(details=[{"field": m, "msg": "required"} for m in missing])
            except ValidationFailed as e:
                invalid += 1
                if len(errors) < MAX_ERRORS:
                    errors.append({"row": n, "errors": e.details})
                continue
            values["id"] = uuid.uuid5(task_id, str(n))
            valid.append(values)
        if invalid and inp["mode"] == "all_or_nothing":
            raise ValidationFailed(
                "Import rejected: some rows are invalid; nothing was imported.",
                details={"invalid_rows": invalid, "errors": errors},
            )

        table = await self._schema.sa_table(meta)
        tenant = await self._tenants.get(tenant_id)
        dedicated = tenant is not None and tenant.isolation_mode == IsolationMode.dedicated
        inserted = 0
        async with self._tenant_db.session(p, dedicated=dedicated) as s:
            for i in range(0, len(valid), 500):
                chunk = [
                    {**v, "tenant_id": tenant_id, "created_by": user_id, "updated_by": user_id}
                    for v in valid[i : i + 500]
                ]
                res = await s.execute(
                    pg_insert(table).values(chunk).on_conflict_do_nothing(index_elements=["id"])
                )
                inserted += res.rowcount or 0  # type: ignore[attr-defined]
            await record(
                s,
                AuditContext(p.kind, p.actor_id, None, None),
                tenant_id=tenant_id,
                action="import.apply",
                resource=meta.resource,
                resource_id=str(task_id),
                detail={"rows": len(valid), "inserted": inserted, "invalid": invalid},
            )
            await s.commit()
        return {
            "table": meta.name,
            "valid_rows": len(valid),
            "inserted": inserted,
            "already_present": len(valid) - inserted,
            "invalid_rows": invalid,
            "errors": errors,
        }

    async def _close(
        self,
        tenant_id: uuid.UUID,
        task_id: uuid.UUID,
        status: TaskStatus,
        result: dict[str, Any] | None,
        error: str | None,
        details: Any,
    ) -> None:
        async with self._db.tenant_session(tenant_id) as s:
            task = (
                await s.execute(select(SystemTask).where(SystemTask.id == task_id).with_for_update())
            ).scalar_one()
            task.status, task.finished_at = status, datetime.now(UTC)
            task.result = (
                result if result is not None else ({"details": details} if details is not None else None)
            )
            task.error = error[:1000] if error else None
            await emit(
                s,
                tenant_id,
                "import.completed" if status == TaskStatus.succeeded else "import.failed",
                {
                    "task_id": task_id,
                    "table": task.input.get("table"),
                    "inserted": (result or {}).get("inserted", 0),
                },
            )
            await s.commit()

    async def get_task(self, p: Principal, task_id: uuid.UUID) -> SystemTask:
        async with self._db.tenant_session(p.tenant_id, user_id=p.user_id) as s:
            task = (await s.execute(select(SystemTask).where(SystemTask.id == task_id))).scalar_one_or_none()
        # Tasks are visible to their creator and to tenant admins; otherwise indistinguishable from absent.
        if task is None or not (task.created_by == p.user_id or p.perms.is_admin):
            raise NotFound()
        return task
